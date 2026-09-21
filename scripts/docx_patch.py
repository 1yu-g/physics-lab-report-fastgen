#!/usr/bin/env python3
"""Apply small text, cell, or embedded-image edits without rebuilding a DOCX."""
from __future__ import annotations

import argparse
import json
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document

import fastgen


def _replace_paragraph_text(paragraph, old: str, new: str):
    if old not in paragraph.text:
        return False
    replacement = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = replacement
    return True


def _image_targets(doc: Document, image_edits, base: Path):
    targets = {}
    for edit in image_edits:
        index = int(edit["index"])
        if index < 0 or index >= len(doc.inline_shapes):
            raise IndexError(f"Image index out of range: {index}")
        replacement = Path(edit["path"])
        replacement = replacement if replacement.is_absolute() else base / replacement
        replacement = replacement.resolve()
        if not replacement.is_file():
            raise FileNotFoundError(replacement)
        shape = doc.inline_shapes[index]
        relation_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        image_part = doc.part.related_parts[relation_id]
        member = str(image_part.partname).lstrip("/")
        if Path(member).suffix.lower() != replacement.suffix.lower():
            raise ValueError("Replacement image must use the same file extension as the embedded image.")
        targets[member] = replacement
    return targets


def _replace_zip_members(docx_path: Path, targets):
    if not targets:
        return
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=docx_path.parent) as stream:
        temporary = Path(stream.name)
    try:
        with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temporary, "w") as destination:
            names = set(source.namelist())
            missing = set(targets) - names
            if missing:
                raise ValueError(f"Embedded image members are missing: {sorted(missing)}")
            for info in source.infolist():
                payload = targets[info.filename].read_bytes() if info.filename in targets else source.read(info.filename)
                destination.writestr(info, payload)
        temporary.replace(docx_path)
    finally:
        temporary.unlink(missing_ok=True)


def apply(spec_path: Path):
    spec_path = spec_path.expanduser().resolve()
    spec = fastgen.load(spec_path)
    base = spec_path.parent
    source = Path(spec["source"])
    output = Path(spec["output"])
    source = (source if source.is_absolute() else base / source).resolve()
    output = (output if output.is_absolute() else base / output).resolve()
    if source == output:
        raise ValueError("Output must differ from the source DOCX.")
    if not source.is_file() or source.suffix.lower() != ".docx":
        raise ValueError(f"Source must be an existing DOCX: {source}")
    source_hash = fastgen.sha256(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    doc = Document(output)
    before = {"tables": len(doc.tables), "images": len(doc.inline_shapes)}
    text_changes = 0
    for edit in spec.get("paragraphs", []):
        matches = 0
        for paragraph in fastgen.paragraphs(doc):
            if _replace_paragraph_text(paragraph, str(edit["find"]), str(edit["replace"])):
                matches += 1
                if not edit.get("all", False):
                    break
        if not matches:
            raise ValueError(f"Paragraph text was not found: {edit['find']}")
        text_changes += matches
    cell_changes = 0
    for edit in spec.get("table_cells", []):
        table = doc.tables[int(edit["table"])]
        table.cell(int(edit["row"]), int(edit["column"])).text = str(edit["text"])
        cell_changes += 1
    image_targets = _image_targets(doc, spec.get("images", []), base)
    doc.save(output)
    _replace_zip_members(output, image_targets)
    if fastgen.sha256(source) != source_hash:
        raise RuntimeError("Source DOCX changed during patching.")
    source_qa = fastgen.check(source)
    qa = fastgen.check(output)
    if qa["tables"] != before["tables"] or qa["inline_images"] != before["images"]:
        raise RuntimeError("Targeted patch changed the document structure.")
    existing_placeholders = set(source_qa.get("unresolved_placeholders", []))
    new_placeholders = set(qa.get("unresolved_placeholders", [])) - existing_placeholders
    if qa.get("package_error") or new_placeholders:
        raise RuntimeError("Patched DOCX is invalid or introduced unresolved placeholders.")
    manifest = {
        "status": "pass", "source": str(source), "source_sha256": source_hash,
        "output": str(output), "output_sha256": fastgen.sha256(output),
        "paragraph_changes": text_changes, "table_cell_changes": cell_changes,
        "image_changes": len(image_targets), "qa": qa,
    }
    manifest_path = output.with_suffix(".patch.json")
    fastgen.dump(manifest_path, manifest)
    return {**manifest, "manifest": str(manifest_path)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", type=Path, required=True)
    args = parser.parse_args()
    try:
        print(json.dumps(apply(args.spec), ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f"docx_patch: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
