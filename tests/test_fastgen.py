import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
MODULE_SPEC = importlib.util.spec_from_file_location("fastgen", ROOT / "scripts" / "fastgen.py")
fastgen = importlib.util.module_from_spec(MODULE_SPEC)
assert MODULE_SPEC.loader is not None
MODULE_SPEC.loader.exec_module(fastgen)


class FastGenIntegrationTest(unittest.TestCase):
    def test_template_copy_build_and_qa(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.docx"
            output = root / "report.docx"
            spec_path = root / "report.json"
            image_path = root / "principle.png"

            doc = Document()
            doc.add_paragraph("姓名：{{姓名}}")
            for title in ("一、实验目的", "二、实验原理", "三、实验仪器", "四、实验步骤"):
                doc.add_paragraph(title, style="Heading 1")
            doc.save(template)
            original_hash = fastgen.sha256(template)

            Image.new("RGB", (320, 180), "white").save(image_path)
            spec = {
                "metadata": {"姓名": "测试用户"},
                "sections": [
                    {
                        "title": "一、实验目的",
                        "blocks": [{"type": "paragraph", "text": "测量物理量。"}],
                    },
                    {
                        "title": "二、实验原理",
                        "blocks": [
                            {"type": "formula", "text": "U_H = K_H I B"},
                            {"type": "figure", "path": "principle.png", "caption": "原理图"},
                        ],
                    },
                    {
                        "title": "三、实验仪器",
                        "blocks": [{"type": "paragraph", "text": "测试仪器。"}],
                    },
                    {
                        "title": "四、实验步骤",
                        "blocks": [
                            {
                                "type": "table",
                                "caption": "数据表",
                                "columns": ["序号", "读数"],
                                "rows": [["1", ""]],
                            }
                        ],
                    },
                ],
            }
            spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
            report = fastgen.build(spec_path, output, template)

            self.assertEqual(report["status"], "pass")
            self.assertEqual(report["inline_images"], 1)
            self.assertEqual(report["native_math_objects"], 1)
            self.assertEqual(len(Document(output).element.xpath(".//m:sSub")), 2)
            self.assertEqual(report["table_blocks"], 1)
            self.assertEqual(fastgen.sha256(template), original_hash)
            self.assertTrue(output.with_suffix(".qa.json").is_file())
            built = Document(output)
            text = "\n".join(p.text for p in fastgen.paragraphs(built))
            self.assertIn("测试用户", text)
            self.assertIn("测量物理量", text)


if __name__ == "__main__":
    unittest.main()
