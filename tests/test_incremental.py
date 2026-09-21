import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import docx_patch
import fastgen
import template_profile
import workflow


class IncrementalWorkflowTest(unittest.TestCase):
    def test_profile_validation_and_preflight(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            template = root / "template.docx"
            doc = Document()
            doc.add_paragraph("姓名：{{姓名}}")
            doc.add_paragraph("一、实验目的", style="Heading 1")
            doc.add_table(rows=2, cols=2)
            doc.save(template)
            profile = root / "profile.json"
            created = template_profile.create(template, profile)
            self.assertEqual(created["status"], "pass")
            self.assertEqual(template_profile.check(profile, template)["tables"], 1)
            work = root / "work"
            prepared = workflow.prepare(
                work, [], template, [], [], ["一、实验目的"], "封面不动", profile
            )
            self.assertEqual(prepared["files"], 1)
            changed = Document(template)
            changed.add_paragraph("二、实验原理")
            changed.save(template)
            with self.assertRaisesRegex(ValueError, "changed after profiling"):
                template_profile.check(profile, template)
        self.assertEqual(workflow.preflight("core")["status"], "pass")

    def test_targeted_docx_patch_preserves_structure_and_source(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            first = root / "first.jpg"
            second = root / "second.jpg"
            replacement = root / "replacement.jpg"
            Image.new("RGB", (120, 80), "red").save(first)
            Image.new("RGB", (120, 80), "green").save(second)
            Image.new("RGB", (120, 80), "blue").save(replacement)
            source = root / "source.docx"
            output = root / "output.docx"
            doc = Document()
            doc.add_paragraph("霍尔电压旧值")
            table = doc.add_table(rows=2, cols=2)
            table.cell(1, 1).text = "旧数据"
            doc.add_picture(str(first), width=Inches(1))
            doc.add_picture(str(second), width=Inches(1))
            doc.save(source)
            source_hash = fastgen.sha256(source)
            spec = root / "patch.json"
            fastgen.dump(spec, {
                "source": str(source), "output": str(output),
                "paragraphs": [{"find": "旧值", "replace": "新值"}],
                "table_cells": [{"table": 0, "row": 1, "column": 1, "text": "新数据"}],
                "images": [{"index": 1, "path": str(replacement)}],
            })
            result = docx_patch.apply(spec)
            self.assertEqual(result["status"], "pass")
            self.assertEqual(fastgen.sha256(source), source_hash)
            patched = Document(output)
            self.assertIn("霍尔电压新值", "\n".join(p.text for p in patched.paragraphs))
            self.assertEqual(patched.tables[0].cell(1, 1).text, "新数据")
            self.assertEqual(len(patched.inline_shapes), 2)
            with zipfile.ZipFile(output) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
                self.assertEqual(len(media), 2)


if __name__ == "__main__":
    unittest.main()
