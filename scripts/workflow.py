#!/usr/bin/env python3
"""Inventory inputs, build a DOCX, render previews, and gate delivery on review."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import platform
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader

import fastgen

HEAD = re.compile(r"^(?:[一二三四五六七八九十]+|\d+)[、.．]\s*\S+")
FIGURE = re.compile(r"(?:图|Fig\.?)\s*\d+[\w.-]*", re.IGNORECASE)
TABLE = re.compile(r"(?:表|Table)\s*\d+[\w.-]*", re.IGNORECASE)
PLACEHOLDER = re.compile(r"\{\{[^{}]+\}\}")
EXTRACT_CACHE_VERSION = "1"


def _cache_key(path: Path):
    return f"{fastgen.sha256(path)}-{path.suffix.lower().lstrip('.') or 'file'}"


def document_text(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages), len(pages), sum(len(page.images) for page in reader.pages)
    if suffix == ".docx":
        doc = Document(path)
        content = [p.text for p in fastgen.paragraphs(doc)]
        return "\n".join(content), None, len(doc.inline_shapes)
    if suffix in {".txt", ".md", ".csv", ".tsv"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return path.read_text(encoding=encoding), None, 0
            except UnicodeDecodeError:
                pass
        raise ValueError(f"Cannot decode source text: {path}")
    return "", None, 0


def record(path: Path, role: str, text_dir: Path, index: int, cache_dir: Path | None = None,
           profile_payload=None):
    path = path.expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    digest = fastgen.sha256(path)
    item = {
        "role": role, "path": str(path), "name": path.name,
        "size_bytes": path.stat().st_size, "sha256": digest,
    }
    if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        with Image.open(path) as image:
            item["dimensions"] = list(image.size)
        item["visual_review_required"] = True
        item["cache_hit"] = False
    else:
        cache_dir = (cache_dir or text_dir / ".cache").resolve()
        cache_dir.mkdir(parents=True, exist_ok=True)
        key = _cache_key(path)
        metadata_path = cache_dir / f"{key}.json"
        cached_text = cache_dir / f"{key}.txt"
        cached = fastgen.load(metadata_path) if metadata_path.is_file() else None
        if profile_payload:
            content = "\n".join(anchor["text"] for anchor in profile_payload.get("anchors", []))
            payload = {
                "pages": None,
                "embedded_images": profile_payload.get("inline_images", 0),
                "characters": len(content),
                "figure_mentions": sorted(set(FIGURE.findall(content))),
                "table_mentions": sorted(set(TABLE.findall(content))),
            }
            item["cache_hit"] = True
            item["profile_hit"] = True
        elif (cached and cached.get("version") == EXTRACT_CACHE_VERSION
                and cached.get("sha256") == digest
                and (not cached.get("has_text") or cached_text.is_file())):
            content = cached_text.read_text(encoding="utf-8") if cached.get("has_text") else ""
            payload = cached["payload"]
            item["cache_hit"] = True
        else:
            content, page_count, image_count = document_text(path)
            payload = {
                "pages": page_count, "embedded_images": image_count,
                "characters": len(content),
                "figure_mentions": sorted(set(FIGURE.findall(content))),
                "table_mentions": sorted(set(TABLE.findall(content))),
            }
            if content:
                cached_text.write_text(content, encoding="utf-8")
            fastgen.dump(metadata_path, {
                "version": EXTRACT_CACHE_VERSION, "sha256": digest,
                "has_text": bool(content), "payload": payload,
            })
            item["cache_hit"] = False
        item.update(payload)
        if content:
            text_dir.mkdir(parents=True, exist_ok=True)
            text_path = text_dir / f"{index:02d}-{role}.txt"
            if not text_path.is_file() or text_path.read_text(encoding="utf-8") != content:
                text_path.write_text(content, encoding="utf-8")
            item["extracted_text"] = str(text_path.resolve())
        if role == "template" and path.suffix.lower() == ".docx":
            item["headings"] = [line.strip() for line in content.splitlines() if HEAD.match(line.strip())]
            item["placeholders"] = sorted(set(PLACEHOLDER.findall(content)))
    return item


def prepare(workdir: Path, guides, template, images, data, sections, scope, profile=None):
    workdir = workdir.resolve()
    workdir.mkdir(parents=True, exist_ok=True)
    files = []
    paths = ([("guide", p) for p in guides] +
             ([("template", template)] if template else []) +
             [("image", p) for p in images] +
             [("data", p) for p in data])
    if not paths:
        raise ValueError("Provide at least one guide, template, image, or data file.")
    if template and Path(template).suffix.lower() != ".docx":
        raise ValueError("The template must be DOCX; convert old DOC files before prepare.")
    profile_payload = None
    if profile:
        if not template:
            raise ValueError("A template profile requires --template.")
        import template_profile
        template_profile.check(Path(profile), Path(template))
        profile_payload = fastgen.load(profile)
    for index, (role, path) in enumerate(paths, 1):
        files.append(record(Path(path), role, workdir / "source-text", index,
                            workdir / ".cache" / "extract",
                            profile_payload if role == "template" else None))
    chosen = list(sections)
    if not chosen:
        template_item = next((item for item in files if item["role"] == "template"), {})
        chosen = template_item.get("headings", [])
    inventory = {
        "scope": scope, "requested_sections": list(sections),
        "template_profile": str(Path(profile).resolve()) if profile else None,
        "template_profile_sha256": fastgen.sha256(profile) if profile else None,
        "files": files,
        "needs_human_or_agent_review": [
            "Confirm requested sections and draft content against the guide.",
            "Select only relevant, clear images and verify they have no watermark.",
            "Verify formula symbols, units, tables, and all real measurements.",
        ],
    }
    inventory_path = workdir / "inventory.json"
    fastgen.dump(inventory_path, inventory)
    spec_path = workdir / "report.json"
    if not spec_path.exists():
        fastgen.dump(spec_path, {
            "title": "大学物理实验报告",
            "metadata": {},
            "sections": [
                {"title": title, "anchor": title, "blocks": []} for title in chosen
            ],
        })
    return {"inventory": str(inventory_path), "draft_spec": str(spec_path),
            "files": len(files), "sections": len(chosen),
            "cache_hits": sum(bool(item.get("cache_hit")) for item in files)}


def validate_spec(spec: dict, base: Path, required_sections):
    sections = spec.get("sections", [])
    if not sections:
        raise ValueError("Report JSON has no sections; complete the draft first.")
    names = [str(section.get("title", "")).strip() for section in sections]
    if any(not name for name in names) or len(set(names)) != len(names):
        raise ValueError("Section titles must be nonempty and unique.")
    if required_sections and names != required_sections:
        raise ValueError("Report sections do not match the requested scope and order.")
    counts = {"figure": 0, "table": 0, "formula": 0}
    for section in sections:
        blocks = section.get("blocks", [])
        if not blocks:
            raise ValueError(f"Section has no content: {section['title']}")
        for block in blocks:
            kind = block.get("type", "paragraph")
            if kind == "paragraph":
                content = block.get("text") or "".join(x.get("text", "") for x in block.get("segments", []))
                if not str(content).strip():
                    raise ValueError(f"Empty paragraph in {section['title']}")
            elif kind == "formula":
                if not str(block.get("text", "")).strip():
                    raise ValueError(f"Empty formula in {section['title']}")
                counts[kind] += 1
            elif kind == "figure":
                path = Path(block.get("path", ""))
                path = path if path.is_absolute() else base / path
                if not path.is_file() or not str(block.get("caption", "")).strip():
                    raise ValueError(f"Figure needs an existing file and caption: {section['title']}")
                counts[kind] += 1
            elif kind == "table":
                if not block.get("columns") or not str(block.get("caption", "")).strip():
                    raise ValueError(f"Table needs columns and caption: {section['title']}")
                counts[kind] += 1
            elif kind != "page_break":
                raise ValueError(f"Unsupported block: {kind}")
    return counts


def find_soffice():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for candidate in (
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
    ):
        if candidate.is_file():
            return str(candidate)
    return None


DEPENDENCY_GROUPS = {
    "core": {"docx": "python-docx", "PIL": "Pillow", "pypdf": "pypdf"},
    "analysis": {
        "pandas": "pandas", "scipy": "scipy", "sympy": "sympy",
        "pint": "pint", "uncertainties": "uncertainties", "matplotlib": "matplotlib",
    },
    "ocr": {"img2table": "img2table", "rapidocr": "img2table[rapidocr]"},
}


def preflight(mode="core"):
    modes = {
        "core": ("core",),
        "analysis": ("core", "analysis"),
        "ocr": ("core", "ocr"),
        "full": tuple(DEPENDENCY_GROUPS),
    }
    if mode not in modes:
        raise ValueError(f"Unknown preflight mode: {mode}")
    groups = {}
    missing = []
    for group in modes[mode]:
        status = {}
        for module, package in DEPENDENCY_GROUPS[group].items():
            available = importlib.util.find_spec(module) is not None
            status[module] = {"available": available, "package": package}
            if not available:
                missing.append(package)
        groups[group] = status
    renderers = {
        "libreoffice": find_soffice(),
        "pdftoppm": shutil.which("pdftoppm"),
    }
    return {
        "status": "pass" if not missing else "needs-install",
        "mode": mode,
        "python": sys.version.split()[0],
        "platform": platform.platform(),
        "groups": groups,
        "missing_packages": sorted(set(missing)),
        "renderers": renderers,
        "render_ready": bool(renderers["libreoffice"] and renderers["pdftoppm"]),
    }


def render(docx_path: Path, preview_dir: Path, pdf_path=None, renderer="auto"):
    preview_dir.mkdir(parents=True, exist_ok=True)
    if pdf_path:
        pdf_path = Path(pdf_path).resolve()
        if not pdf_path.is_file():
            raise FileNotFoundError(pdf_path)
    elif renderer != "none":
        soffice = find_soffice()
        if not soffice:
            return None, []
        profile = (preview_dir / "libreoffice-profile").resolve().as_uri()
        subprocess.run([
            soffice, "--headless", f"-env:UserInstallation={profile}",
            "--convert-to", "pdf", "--outdir", str(preview_dir), str(docx_path)
        ], check=True, capture_output=True, text=True, timeout=120)
        pdf_path = preview_dir / (docx_path.stem + ".pdf")
        if not pdf_path.is_file():
            raise RuntimeError("LibreOffice did not create a PDF.")
    else:
        return None, []
    if pdf_path.stat().st_mtime < docx_path.stat().st_mtime - 2:
        raise ValueError("The supplied PDF predates the DOCX; export the latest report again.")
    count = len(PdfReader(pdf_path).pages)
    if count == 0:
        raise ValueError("Rendered PDF has no pages.")
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        return pdf_path, []
    for old_preview in preview_dir.glob("page-*.png"):
        old_preview.unlink()
    prefix = preview_dir / "page"
    subprocess.run([
        pdftoppm, "-f", "1", "-l", str(count), "-r", "120",
        "-png", str(pdf_path), str(prefix)
    ], check=True, capture_output=True, text=True, timeout=120)
    previews = sorted(preview_dir.glob("page-*.png"), key=lambda path: int(path.stem.split("-")[-1]))
    if len(previews) != count:
        raise RuntimeError(f"Expected {count} page previews, got {len(previews)}.")
    return pdf_path, previews


RESULT_TOKEN = re.compile(r"\{\{result\.([A-Za-z_][A-Za-z0-9_.]*)(?::([^{}]+))?\}\}")


def resolve_results(spec, analysis):
    def lookup(reference):
        value = analysis["results"]
        for part in reference.split("."):
            if not isinstance(value, dict) or part not in value:
                raise ValueError(f"Unknown analysis result: {reference}")
            value = value[part]
        if isinstance(value, (dict, list)):
            raise ValueError(f"Analysis result is not scalar: {reference}")
        return value

    def replace(value):
        if isinstance(value, dict):
            return {key: replace(item) for key, item in value.items()}
        if isinstance(value, list):
            return [replace(item) for item in value]
        if isinstance(value, str):
            def match_value(match):
                result = lookup(match.group(1))
                fmt = match.group(2)
                return format(result, fmt) if fmt else str(result)
            return RESULT_TOKEN.sub(match_value, value)
        return value
    return replace(spec)


def _build_fingerprint(spec_path: Path, spec: dict, template: Path | None):
    files = {}
    for section in spec.get("sections", []):
        for block in section.get("blocks", []):
            if block.get("type") != "figure" or not block.get("path"):
                continue
            path = Path(block["path"])
            path = path if path.is_absolute() else spec_path.parent / path
            path = path.resolve()
            if not path.is_file():
                raise FileNotFoundError(path)
            files[str(path)] = fastgen.sha256(path)
    payload = {
        "builder_version": fastgen.VERSION,
        "spec_sha256": fastgen.sha256(spec_path),
        "template_sha256": fastgen.sha256(template) if template else None,
        "referenced_files": files,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(raw).hexdigest(), payload


def run(workdir: Path, spec_path=None, template=None, output=None, pdf=None, renderer="auto"):
    workdir = workdir.resolve()
    inventory_path = workdir / "inventory.json"
    inventory = fastgen.load(inventory_path) if inventory_path.is_file() else {}
    spec_path = Path(spec_path or workdir / "report.json").resolve()
    spec = fastgen.load(spec_path)
    plan_ref, manifest_ref = spec.get("analysis_plan"), spec.get("analysis_manifest")
    if plan_ref and manifest_ref:
        raise ValueError("Use analysis_plan or analysis_manifest, not both.")
    analysis_path = None
    if plan_ref or manifest_ref:
        import analyze_data
        if plan_ref:
            plan_path = Path(plan_ref)
            plan_path = (plan_path if plan_path.is_absolute()
                         else spec_path.parent / plan_path).resolve()
            analysis_path = Path(analyze_data.analyze(plan_path, workdir / "analysis")["analysis"])
        else:
            analysis_path = Path(manifest_ref)
            analysis_path = (analysis_path if analysis_path.is_absolute()
                             else spec_path.parent / analysis_path).resolve()
        analysis = analyze_data.check_analysis(analysis_path)
        spec = resolve_results(spec, analysis)
        for section_index, section in enumerate(spec.get("sections", []), 1):
            for block_index, block in enumerate(section.get("blocks", []), 1):
                if block.get("type") != "figure":
                    continue
                if block.get("fit_id"):
                    if block.get("path"):
                        raise ValueError("Fit figure uses fit_id instead of path.")
                    import figure_tools
                    figure_path = workdir / "figures" / f"fit-{section_index}-{block_index}.png"
                    figure_tools.plot_fit(analysis_path, str(block["fit_id"]),
                                          figure_path, block.get("plot_title", ""))
                    block["path"] = str(figure_path.resolve())
                else:
                    image_path = Path(block.get("path", ""))
                    if not image_path.is_absolute():
                        block["path"] = str((spec_path.parent / image_path).resolve())
        resolved_spec = workdir / "resolved-report.json"
        fastgen.dump(resolved_spec, spec)
        build_spec_path = resolved_spec
    else:
        if RESULT_TOKEN.search(json.dumps(spec, ensure_ascii=False)):
            raise ValueError("Result placeholders need an analysis_plan or analysis_manifest.")
        build_spec_path = spec_path
    counts = validate_spec(spec, build_spec_path.parent, inventory.get("requested_sections", []))
    if template is None:
        template_item = next((x for x in inventory.get("files", []) if x["role"] == "template"), None)
        template = Path(template_item["path"]) if template_item else None
    output = Path(output or workdir / "output" / "report.docx").resolve()
    template_hash = fastgen.sha256(template) if template else None
    template_images = len(Document(template).inline_shapes) if template else 0
    template_math = len(Document(template).element.xpath(".//m:oMath")) if template else 0
    fingerprint, fingerprint_payload = _build_fingerprint(build_spec_path, spec, template)
    build_cache = workdir / ".cache" / "build.json"
    cached = fastgen.load(build_cache) if build_cache.is_file() else {}
    qa_path = output.with_suffix(".qa.json")
    cache_hit = bool(
        cached.get("fingerprint") == fingerprint
        and output.is_file() and qa_path.is_file()
        and cached.get("output_sha256") == fastgen.sha256(output)
    )
    if cache_hit:
        qa = fastgen.load(qa_path)
    else:
        qa = fastgen.build(build_spec_path, output, template)
        if template and fastgen.sha256(template) != template_hash:
            raise RuntimeError("Source template changed during build.")
        for key, actual_key, baseline in (("figure", "inline_images", template_images),
                                          ("formula", "native_math_objects", template_math)):
            actual = qa[actual_key] - baseline
            if actual != counts[key]:
                qa["status"] = "needs-fix"
                qa.setdefault("count_mismatch", {})[key] = {"expected": counts[key], "actual": actual}
        if qa["table_blocks"] != counts["table"]:
            qa["status"] = "needs-fix"
        if analysis_path:
            qa["analysis_manifest"] = str(analysis_path)
            qa["analysis_sha256"] = fastgen.sha256(analysis_path)
        fastgen.dump(qa_path, qa)
        fastgen.dump(build_cache, {
            "fingerprint": fingerprint, "inputs": fingerprint_payload,
            "output": str(output), "output_sha256": fastgen.sha256(output),
        })
    if qa["status"] != "pass":
        raise RuntimeError(f"Structural QA failed: {qa_path}")
    result = _write_review(workdir, output, pdf, renderer)
    result["build_cache_hit"] = cache_hit
    return result


def _write_review(workdir: Path, docx: Path, pdf, renderer):
    review_path = workdir / "review.json"
    old_review = fastgen.load(review_path) if review_path.is_file() else {}
    docx_hash = fastgen.sha256(docx)
    old_pages = {}
    for page in old_review.get("pages", []):
        preview = Path(page.get("preview", ""))
        digest = page.get("preview_sha256")
        if not digest and preview.is_file():
            digest = fastgen.sha256(preview)
        if digest:
            old_pages[int(page["page"])] = {"sha256": digest, "status": page.get("status", "pending")}
    cached_pages_valid = all(
        Path(page.get("preview", "")).is_file()
        and page.get("preview_sha256") == fastgen.sha256(page["preview"])
        for page in old_review.get("pages", [])
    )
    if (pdf is None and old_review.get("docx_sha256") == docx_hash
            and old_review.get("pages") and cached_pages_valid):
        return {
            "status": old_review["status"], "docx": str(docx),
            "qa": str(docx.with_suffix(".qa.json")), "review": str(review_path),
            "pages": len(old_review["pages"]), "changed_pages": [],
            "render_cache_hit": True,
        }
    pdf_path, previews = render(docx, workdir / "preview", pdf, renderer)
    pages = []
    changed_pages = []
    for index, path in enumerate(previews, 1):
        digest = fastgen.sha256(path)
        old = old_pages.get(index)
        unchanged = bool(old and old["sha256"] == digest)
        status = "pass" if unchanged and old.get("status") == "pass" else "pending"
        if not unchanged:
            changed_pages.append(index)
        pages.append({
            "page": index, "preview": str(path), "preview_sha256": digest,
            "changed": not unchanged, "status": status, "notes": "",
        })
    review = {
        "status": "needs-visual-review" if previews else "render-unavailable",
        "docx": str(docx), "docx_sha256": docx_hash,
        "pdf": str(pdf_path) if pdf_path else None,
        "pdf_sha256": fastgen.sha256(pdf_path) if pdf_path else None,
        "criteria": ["scope", "cover", "no_watermark", "page_layout", "figures",
                     "tables", "formula_symbols", "units", "captions", "data_provenance"],
        "pages": pages,
        "changed_pages": changed_pages,
    }
    fastgen.dump(review_path, review)
    return {"status": review["status"], "docx": str(docx),
            "qa": str(docx.with_suffix(".qa.json")),
            "review": str(review_path), "pages": len(previews),
            "changed_pages": changed_pages, "render_cache_hit": False}


def _checked_report(workdir: Path):
    review = fastgen.load(workdir / "review.json")
    docx = Path(review["docx"])
    qa = fastgen.load(docx.with_suffix(".qa.json"))
    if qa["status"] != "pass" or fastgen.sha256(docx) != review["docx_sha256"]:
        raise ValueError("DOCX changed after QA or structural QA did not pass; run again.")
    if qa.get("analysis_manifest"):
        import analyze_data
        analysis_path = Path(qa["analysis_manifest"])
        if fastgen.sha256(analysis_path) != qa["analysis_sha256"]:
            raise ValueError("Analysis changed after report generation; run again.")
        analyze_data.check_analysis(analysis_path)
    return review, docx


def preview(workdir: Path, pdf: Path):
    workdir = workdir.resolve()
    _, docx = _checked_report(workdir)
    return _write_review(workdir, docx, pdf, "none")


def finalize(workdir: Path):
    workdir = workdir.resolve()
    review, docx = _checked_report(workdir)
    pages = review.get("pages", [])
    if (not pages or
        [page.get("page") for page in pages] != list(range(1, len(pages) + 1)) or
        any(page.get("status") != "pass"
            or not Path(page.get("preview", "")).is_file()
            or page.get("preview_sha256") != fastgen.sha256(page["preview"])
            for page in pages)):
        raise ValueError("Inspect every rendered page and set each page status to pass before finalizing.")
    result = {"status": "pass", "docx": str(docx), "pdf": review.get("pdf"),
              "pages_reviewed": len(pages), "qa": str(docx.with_suffix(".qa.json"))}
    fastgen.dump(workdir / "delivery.json", result)
    return result


def cli():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("prepare", help="Inventory materials and create a blank report JSON")
    p.add_argument("--workdir", type=Path, required=True)
    p.add_argument("--guide", action="append", type=Path, default=[])
    p.add_argument("--template", type=Path)
    p.add_argument("--image", action="append", type=Path, default=[])
    p.add_argument("--data", action="append", type=Path, default=[])
    p.add_argument("--section", action="append", default=[])
    p.add_argument("--scope", default="")
    p.add_argument("--profile", type=Path,
                   help="Validate the template against a reusable template profile")
    p = sub.add_parser("run", help="Build, structurally check, and render page previews")
    p.add_argument("--workdir", type=Path, required=True)
    p.add_argument("--spec", type=Path)
    p.add_argument("--template", type=Path)
    p.add_argument("--output", type=Path)
    p.add_argument("--pdf", type=Path, help="Use a PDF already exported from this DOCX")
    p.add_argument("--renderer", choices=("auto", "none"), default="auto")
    p = sub.add_parser("preview", help="Render an exported PDF without rebuilding the DOCX")
    p.add_argument("--workdir", type=Path, required=True)
    p.add_argument("--pdf", type=Path, required=True)
    p = sub.add_parser("finalize", help="Verify page-by-page review and mark delivery ready")
    p.add_argument("--workdir", type=Path, required=True)
    p = sub.add_parser("preflight", help="Check the runtime before starting a report")
    p.add_argument("--mode", choices=("core", "analysis", "ocr", "full"), default="core")
    return parser.parse_args()


def main():
    args = cli()
    try:
        if args.command == "prepare":
            result = prepare(args.workdir, args.guide, args.template, args.image,
                             args.data, args.section, args.scope, args.profile)
        elif args.command == "run":
            result = run(args.workdir, args.spec, args.template, args.output, args.pdf, args.renderer)
        elif args.command == "preview":
            result = preview(args.workdir, args.pdf)
        elif args.command == "preflight":
            result = preflight(args.mode)
        else:
            result = finalize(args.workdir)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f"workflow: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
