#!/usr/bin/env python3
"""Extract common lab materials into cached text, tables, images, and provenance JSON."""
from __future__ import annotations

import argparse
import csv
import json
import posixpath
import re
import shutil
import sys
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from docx import Document
from pypdf import PdfReader

import fastgen

VERSION = "1"
SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".csv", ".tsv"}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _decode(path: Path):
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError(f"Cannot decode text source: {path}")


def _write_csv(path: Path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        csv.writer(stream).writerows(rows)
    return {
        "path": str(path), "sha256": fastgen.sha256(path),
        "rows": len(rows), "columns": max((len(row) for row in rows), default=0),
    }


def _copy_zip_assets(archive: zipfile.ZipFile, prefix: str, output_dir: Path):
    assets = []
    for index, name in enumerate(sorted(n for n in archive.namelist() if n.startswith(prefix)), 1):
        if name.endswith("/"):
            continue
        suffix = Path(name).suffix.lower()
        target = output_dir / f"image-{index:03d}{suffix}"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(archive.read(name))
        assets.append({
            "path": str(target), "sha256": fastgen.sha256(target),
            "source_member": name, "kind": "image",
        })
    return assets


def _parse_pdf(source: Path, output_dir: Path):
    reader = PdfReader(source)
    units, assets = [], []
    asset_dir = output_dir / "assets"
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        units.append({"kind": "page", "index": page_number, "text": text})
        for image_number, image in enumerate(page.images, 1):
            name = Path(getattr(image, "name", "image.bin")).name
            suffix = Path(name).suffix or ".bin"
            target = asset_dir / f"page-{page_number:03d}-image-{image_number:03d}{suffix}"
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(image.data)
            assets.append({
                "path": str(target), "sha256": fastgen.sha256(target),
                "page": page_number, "kind": "image",
            })
    return units, [], assets


def _parse_docx(source: Path, output_dir: Path):
    doc = Document(source)
    units = [
        {"kind": "paragraph", "index": index, "text": paragraph.text}
        for index, paragraph in enumerate(fastgen.paragraphs(doc), 1)
        if paragraph.text.strip()
    ]
    tables = []
    for index, table in enumerate(doc.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        item = _write_csv(output_dir / "tables" / f"table-{index:03d}.csv", rows)
        item.update({"kind": "table", "index": index})
        tables.append(item)
    with zipfile.ZipFile(source) as archive:
        assets = _copy_zip_assets(archive, "word/media/", output_dir / "assets")
    return units, tables, assets


def _numeric_name(path: str):
    match = re.search(r"(\d+)(?=\.xml$)", path)
    return int(match.group(1)) if match else 0


def _parse_pptx(source: Path, output_dir: Path):
    units, tables = [], []
    with zipfile.ZipFile(source) as archive:
        slides = sorted(
            (name for name in archive.namelist()
             if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=_numeric_name,
        )
        table_index = 0
        for slide_number, name in enumerate(slides, 1):
            root = ET.fromstring(archive.read(name))
            text = "\n".join(
                value.text or "" for value in root.iter()
                if value.tag.endswith("}t") and (value.text or "").strip()
            )
            units.append({"kind": "slide", "index": slide_number, "text": text})
            for table in (node for node in root.iter() if node.tag.endswith("}tbl")):
                table_index += 1
                rows = []
                for row in (node for node in table if node.tag.endswith("}tr")):
                    values = []
                    for cell in (node for node in row if node.tag.endswith("}tc")):
                        values.append("".join(
                            node.text or "" for node in cell.iter() if node.tag.endswith("}t")
                        ))
                    rows.append(values)
                item = _write_csv(
                    output_dir / "tables" / f"table-{table_index:03d}.csv", rows)
                item.update({"kind": "table", "index": table_index, "slide": slide_number})
                tables.append(item)
        assets = _copy_zip_assets(archive, "ppt/media/", output_dir / "assets")
    return units, tables, assets


def _column_index(reference: str):
    letters = "".join(character for character in reference if character.isalpha()).upper()
    value = 0
    for character in letters:
        value = value * 26 + ord(character) - 64
    return max(value - 1, 0)


def _worksheet_rows(root, shared):
    rows = []
    for row in root.findall(f".//{{{MAIN_NS}}}row"):
        values = []
        for cell in row.findall(f"{{{MAIN_NS}}}c"):
            index = _column_index(cell.attrib.get("r", "A1"))
            while len(values) <= index:
                values.append("")
            kind = cell.attrib.get("t")
            if kind == "inlineStr":
                value = "".join(node.text or "" for node in cell.findall(f".//{{{MAIN_NS}}}t"))
            else:
                value_node = cell.find(f"{{{MAIN_NS}}}v")
                value = "" if value_node is None else value_node.text or ""
                if kind == "s" and value:
                    value = shared[int(value)]
                elif kind == "b":
                    value = "TRUE" if value == "1" else "FALSE"
            values[index] = value
        rows.append(values)
    return rows


def _resolve_target(base: str, target: str):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(str(PurePosixPath(base).parent.joinpath(target)))


def _parse_xlsx(source: Path, output_dir: Path):
    units, tables = [], []
    with zipfile.ZipFile(source) as archive:
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = ["".join(node.text or "" for node in item.iter() if node.tag.endswith("}t"))
                      for item in root if item.tag.endswith("}si")]
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            item.attrib["Id"]: _resolve_target("xl/workbook.xml", item.attrib["Target"])
            for item in relationships.findall(f"{{{PKG_REL_NS}}}Relationship")
        }
        sheets = workbook.findall(f".//{{{MAIN_NS}}}sheet")
        for index, sheet in enumerate(sheets, 1):
            name = sheet.attrib.get("name", f"Sheet{index}")
            relation_id = sheet.attrib.get(f"{{{REL_NS}}}id")
            member = targets[relation_id]
            rows = _worksheet_rows(ET.fromstring(archive.read(member)), shared)
            item = _write_csv(output_dir / "tables" / f"sheet-{index:03d}.csv", rows)
            item.update({"kind": "sheet", "index": index, "name": name})
            tables.append(item)
            units.append({
                "kind": "sheet", "index": index, "name": name,
                "text": "\n".join("\t".join(row) for row in rows),
            })
        assets = _copy_zip_assets(archive, "xl/media/", output_dir / "assets")
    return units, tables, assets


def _parse_delimited(source: Path, output_dir: Path):
    delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
    text = _decode(source)
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    item = _write_csv(output_dir / "tables" / "table-001.csv", rows)
    item.update({"kind": "table", "index": 1})
    return [{"kind": "table", "index": 1, "text": text}], [item], []


def _parse(source: Path, output_dir: Path):
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(source, output_dir)
    if suffix == ".docx":
        return _parse_docx(source, output_dir)
    if suffix == ".pptx":
        return _parse_pptx(source, output_dir)
    if suffix == ".xlsx":
        return _parse_xlsx(source, output_dir)
    if suffix in {".csv", ".tsv"}:
        return _parse_delimited(source, output_dir)
    if suffix in {".txt", ".md"}:
        return [{"kind": "document", "index": 1, "text": _decode(source)}], [], []
    if suffix in IMAGE_TYPES:
        target = output_dir / "assets" / source.name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        return [], [], [{"path": str(target), "sha256": fastgen.sha256(target), "kind": "image"}]
    raise ValueError(f"Unsupported material type: {suffix or source.name}")


def _cache_valid(manifest, source: Path):
    if (manifest.get("version") != VERSION
            or manifest.get("source_sha256") != fastgen.sha256(source)):
        return False
    checks = [(manifest.get("text_file"), manifest.get("text_sha256"))]
    checks += [(item.get("path"), item.get("sha256")) for item in manifest.get("tables", [])]
    checks += [(item.get("path"), item.get("sha256")) for item in manifest.get("assets", [])]
    return all(path and digest and Path(path).is_file() and fastgen.sha256(path) == digest
               for path, digest in checks)


def _reset_managed_output(output_dir: Path, manifest_path: Path):
    if not output_dir.exists():
        output_dir.mkdir(parents=True, exist_ok=True)
        return
    if not manifest_path.is_file():
        if any(output_dir.iterdir()):
            raise ValueError("Output directory is not empty and has no FastGen material manifest.")
        return
    manifest = fastgen.load(manifest_path)
    managed = [manifest.get("text_file")]
    managed += [item.get("path") for item in manifest.get("tables", [])]
    managed += [item.get("path") for item in manifest.get("assets", [])]
    managed.append(str(manifest_path))
    for value in managed:
        if not value:
            continue
        path = Path(value).resolve()
        try:
            path.relative_to(output_dir)
        except ValueError as exc:
            raise ValueError("Material manifest points outside its output directory.") from exc
        path.unlink(missing_ok=True)
    for directory in (output_dir / "tables", output_dir / "assets"):
        if directory.is_dir() and not any(directory.iterdir()):
            directory.rmdir()


def ingest(source: Path, output_dir: Path, force=False):
    source = source.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.suffix.lower() not in SUPPORTED | IMAGE_TYPES:
        raise ValueError(f"Unsupported material type: {source.suffix or source.name}")
    manifest_path = output_dir / "material.json"
    if not force and manifest_path.is_file():
        manifest = fastgen.load(manifest_path)
        if _cache_valid(manifest, source):
            return {
                "status": "pass", "manifest": str(manifest_path),
                "text": manifest["text_file"], "cache_hit": True,
                "units": len(manifest["units"]), "tables": len(manifest["tables"]),
                "assets": len(manifest["assets"]),
            }
    _reset_managed_output(output_dir, manifest_path)
    units, tables, assets = _parse(source, output_dir)
    chunks = []
    for unit in units:
        label = unit["kind"].title()
        name = f" {unit.get('name')}" if unit.get("name") else f" {unit.get('index')}"
        chunks.append(f"## {label}{name}\n\n{unit.get('text', '').strip()}")
    text_file = output_dir / "content.md"
    text_file.write_text("\n\n".join(chunks).strip() + "\n", encoding="utf-8")
    manifest = {
        "version": VERSION, "status": "pass",
        "source": str(source), "source_sha256": fastgen.sha256(source),
        "source_type": source.suffix.lower(), "text_file": str(text_file),
        "text_sha256": fastgen.sha256(text_file),
        "units": units, "tables": tables, "assets": assets,
        "needs_ocr": source.suffix.lower() in IMAGE_TYPES or (
            source.suffix.lower() == ".pdf" and sum(len(x.get("text", "")) for x in units) < 40
        ),
    }
    fastgen.dump(manifest_path, manifest)
    return {
        "status": "pass", "manifest": str(manifest_path), "text": str(text_file),
        "cache_hit": False, "units": len(units), "tables": len(tables),
        "assets": len(assets),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    try:
        print(json.dumps(ingest(args.input, args.output_dir, args.force), ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f"material_ingest: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
