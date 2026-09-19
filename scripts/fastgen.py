#!/usr/bin/env python3
"""Fast, deterministic DOCX builder for Chinese physics lab reports."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

VERSION = "0.1.0"
PLACEHOLDER = re.compile(r"\{\{[^{}]+\}\}")


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def dump(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def font(run, name="宋体", size=10.5):
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def configure(doc, cfg, preserve_page=False):
    if not preserve_page:
        normal = doc.styles["Normal"]
        normal.font.name = cfg.get("body_font", "宋体")
        normal.font.size = Pt(cfg.get("body_size_pt", 10.5))
        normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cfg.get("body_font", "宋体"))
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2.5)
            section.left_margin = section.right_margin = Cm(2.6)


def replace_fields(doc, fields):
    tokens = {"{{" + k + "}}": str(v) for k, v in fields.items()}
    if any(token == value for token, value in tokens.items()):
        raise ValueError("A replacement value cannot equal its own placeholder.")
    replaced = []
    for p in paragraphs(doc):
        for token, value in tokens.items():
            while token in p.text:
                start = p.text.index(token)
                end = start + len(token)
                spans, offset = [], 0
                for run in p.runs:
                    spans.append((run, offset, offset + len(run.text)))
                    offset += len(run.text)
                first = next((item for item in spans if item[1] <= start < item[2]), None)
                last = next((item for item in spans if item[1] < end <= item[2]), None)
                if first is None or last is None:
                    raise ValueError(f"Cannot replace split placeholder: {token}")
                left_run, left_start, _ = first
                right_run, right_start, _ = last
                prefix = left_run.text[:start - left_start]
                suffix = right_run.text[end - right_start:]
                if left_run is right_run:
                    left_run.text = prefix + value + suffix
                else:
                    left_run.text = prefix + value
                    clearing = False
                    for run, _, _ in spans:
                        if run is left_run:
                            clearing = True
                        elif run is right_run:
                            run.text = suffix
                            break
                        elif clearing:
                            run.text = ""
                replaced.append(token)
    return sorted(set(replaced))


def find_anchor(doc, text):
    for p in paragraphs(doc):
        if text and text.strip() == p.text.strip():
            return p
    for p in paragraphs(doc):
        if text and text.strip() in p.text.strip():
            return p
    return None


def body_paragraph(doc, block, cfg):
    p = doc.add_paragraph()
    items = block.get("segments") or [{"text": block.get("text", "")}]
    for item in items:
        r = p.add_run(str(item.get("text", "")))
        font(r, item.get("font", cfg.get("body_font", "宋体")), item.get("size_pt", cfg.get("body_size_pt", 10.5)))
        r.bold = bool(item.get("bold"))
        r.italic = bool(item.get("italic"))
        r.font.subscript = bool(item.get("subscript"))
        r.font.superscript = bool(item.get("superscript"))
    pf = p.paragraph_format
    pf.line_spacing = block.get("line_spacing", cfg.get("line_spacing", 1.5))
    if block.get("first_line_indent", True):
        pf.first_line_indent = Pt(cfg.get("body_size_pt", 10.5) * 2)
    pf.space_after = Pt(block.get("space_after_pt", 0))
    pf.alignment = {"left": 0, "center": 1, "right": 2, "justify": 3}.get(block.get("alignment", "justify"), 3)
    return [p._p]


def math_run(value):
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = value
    text.set(qn("xml:space"), "preserve")
    run.append(text)
    return run


def formula(doc, block, cfg):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    mp, om = OxmlElement("m:oMathPara"), OxmlElement("m:oMath")
    expression = str(block.get("text", ""))
    symbol = re.compile(r"([A-Za-zΑ-ω]+)([_^])(?:\{([^{}]+)\}|([A-Za-z0-9Α-ω]+))")
    cursor = 0
    for match in symbol.finditer(expression):
        if match.start() > cursor:
            om.append(math_run(expression[cursor:match.start()]))
        script = OxmlElement("m:sSub" if match.group(2) == "_" else "m:sSup")
        base = OxmlElement("m:e")
        base.append(math_run(match.group(1)))
        part = OxmlElement("m:sub" if match.group(2) == "_" else "m:sup")
        part.append(math_run(match.group(3) or match.group(4)))
        script.extend((base, part))
        om.append(script)
        cursor = match.end()
    if cursor < len(expression):
        om.append(math_run(expression[cursor:]))
    mp.append(om)
    p._p.append(mp)
    if block.get("number"):
        r = p.add_run("    " + str(block["number"]))
        font(r, "Cambria Math", cfg.get("body_size_pt", 10.5))
    return [p._p]


def figure(doc, block, cfg, base, number):
    path = Path(block["path"])
    path = path if path.is_absolute() else base / path
    path = path.resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Figure not found: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(block.get("width_cm", cfg.get("figure_width_cm", 14.5))))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"图{number} {block.get('caption', path.stem)}")
    font(r, cfg.get("caption_font", "宋体"), cfg.get("caption_size_pt", 9))
    return [p._p, c._p]


def repeat_header(row):
    prop = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    prop.append(item)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def cell_text(cell, value, cfg, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(value))
    font(r, cfg.get("table_font", "宋体"), cfg.get("table_size_pt", 9))
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, block, cfg, number):
    columns = block.get("columns", [])
    if not columns:
        raise ValueError("Table block requires columns.")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(f"表{number} {block.get('caption', '')}".rstrip())
    font(r, cfg.get("caption_font", "宋体"), cfg.get("caption_size_pt", 9))
    tbl = doc.add_table(rows=1, cols=len(columns))
    tbl.style = block.get("style", "Table Grid")
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(columns):
        cell_text(tbl.rows[0].cells[i], value, cfg, True)
        pr = tbl.rows[0].cells[i]._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), block.get("header_fill", "D9EAF7"))
        pr.append(shd)
    repeat_header(tbl.rows[0])
    no_split(tbl.rows[0])
    for values in block.get("rows", []):
        row = tbl.add_row()
        no_split(row)
        for i in range(len(columns)):
            cell_text(row.cells[i], values[i] if i < len(values) else "", cfg)
    return [cap._p, tbl._tbl]


def make_blocks(doc, blocks, cfg, base, counters):
    out = []
    for block in blocks:
        kind = block.get("type", "paragraph")
        if kind == "paragraph":
            out += body_paragraph(doc, block, cfg)
        elif kind == "formula":
            out += formula(doc, block, cfg)
        elif kind == "figure":
            counters["figure"] += 1
            out += figure(doc, block, cfg, base, counters["figure"])
        elif kind == "table":
            counters["table"] += 1
            out += table(doc, block, cfg, counters["table"])
        elif kind == "page_break":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            out.append(p._p)
        else:
            raise ValueError(f"Unsupported block type: {kind}")
    return out


def move_after(anchor, elements):
    cursor = anchor
    for element in elements:
        cursor.addnext(element)
        cursor = element


def new_header(doc, spec, cfg):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(spec.get("title", "大学物理实验报告"))
    font(r, cfg.get("heading_font", "黑体"), 18)
    r.bold = True
    if spec.get("metadata"):
        tbl = doc.add_table(rows=0, cols=2)
        tbl.style = "Table Grid"
        for key, value in spec["metadata"].items():
            cells = tbl.add_row().cells
            cell_text(cells[0], key, cfg, True)
            cell_text(cells[1], value, cfg)


def check(path, spec=None):
    path = Path(path)
    package_error, entries = None, 0
    try:
        with zipfile.ZipFile(path) as z:
            entries = len(z.infolist())
            bad = z.testzip()
            package_error = f"Corrupt ZIP entry: {bad}" if bad else None
    except Exception as exc:
        package_error = str(exc)
    doc = Document(path)
    text = "\n".join(p.text for p in paragraphs(doc))
    unresolved = sorted(set(PLACEHOLDER.findall(text)))
    required = [s.get("title", "") for s in (spec or {}).get("sections", [])]
    missing = [x for x in required if x and x not in text]
    status = "pass" if not package_error and not unresolved and not missing else "needs-fix"
    return {
        "status": status,
        "file": str(path.resolve()),
        "size_bytes": path.stat().st_size,
        "zip_entries": entries,
        "package_error": package_error,
        "paragraphs": len(list(paragraphs(doc))),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "native_math_objects": len(doc.element.xpath(".//m:oMath")),
        "unresolved_placeholders": unresolved,
        "missing_sections": missing,
    }


def build(spec_path, output, template=None, qa=None):
    spec_path, output = Path(spec_path), Path(output)
    spec, cfg = load(spec_path), {}
    cfg.update(spec.get("settings", {}))
    output.parent.mkdir(parents=True, exist_ok=True)
    if template:
        template = Path(template)
        if template.resolve() == output.resolve():
            raise ValueError("Output must differ from the source template.")
        shutil.copy2(template, output)
        doc = Document(output)
        normal_font = doc.styles["Normal"].font
        if normal_font.name:
            cfg.setdefault("body_font", normal_font.name)
        if normal_font.size:
            cfg.setdefault("body_size_pt", normal_font.size.pt)
    else:
        doc = Document()
        new_header(doc, spec, cfg)
    configure(doc, cfg, preserve_page=bool(template))
    fields = {**spec.get("fields", {}), **spec.get("metadata", {})}
    replaced = replace_fields(doc, fields)
    counters = {"figure": 0, "table": 0}
    for section in spec.get("sections", []):
        title = str(section.get("title", "")).strip()
        anchor_text = str(section.get("anchor", title)).strip()
        anchor = find_anchor(doc, anchor_text)
        if anchor is None:
            anchor = doc.add_paragraph(title, style=section.get("style", "Heading 1"))
        elif anchor_text.startswith("{{"):
            anchor.text = title
        elements = make_blocks(doc, section.get("blocks", []), cfg, spec_path.resolve().parent, counters)
        move_after(anchor._p, elements)
    doc.save(output)
    report = check(output, spec)
    report.update({
        "template": str(Path(template).resolve()) if template else None,
        "template_sha256": sha256(template) if template else None,
        "source_spec": str(spec_path.resolve()),
        "replaced_fields": replaced,
        "figure_blocks": counters["figure"],
        "table_blocks": counters["table"],
    })
    dump(qa or output.with_suffix(".qa.json"), report)
    return report


def starter():
    return {
        "title": "大学物理实验预习报告",
        "metadata": {"课程名称": "大学物理实验", "实验名称": "示例实验", "姓名": "", "学号": ""},
        "settings": {"body_font": "宋体", "heading_font": "黑体", "body_size_pt": 10.5, "line_spacing": 1.5},
        "sections": [
            {"title": "一、实验目的", "blocks": [{"type": "paragraph", "text": "根据指导书填写实验目的。"}]},
            {"title": "二、实验原理", "blocks": [{"type": "formula", "text": "U_H = K_H I B"}]},
            {"title": "三、实验仪器", "blocks": [{"type": "paragraph", "text": "列出仪器及其用途。"}]},
            {"title": "四、实验步骤", "blocks": [{"type": "table", "caption": "原始数据记录表", "columns": ["序号", "测量量", "单位"], "rows": [["1", "", ""]]}]},
        ],
    }


def cli():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--version", action="version", version=VERSION)
    sub = p.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("init")
    a.add_argument("--output", type=Path, required=True)
    a = sub.add_parser("build")
    a.add_argument("--spec", type=Path, required=True)
    a.add_argument("--template", type=Path)
    a.add_argument("--output", type=Path, required=True)
    a.add_argument("--qa", type=Path)
    a = sub.add_parser("check")
    a.add_argument("--input", type=Path, required=True)
    a.add_argument("--spec", type=Path)
    a.add_argument("--report", type=Path)
    return p.parse_args()


def main():
    args = cli()
    try:
        if args.cmd == "init":
            dump(args.output, starter())
            print(args.output.resolve())
            return 0
        if args.cmd == "build":
            result = build(args.spec, args.output, args.template, args.qa)
        else:
            result = check(args.input, load(args.spec) if args.spec else None)
            dump(args.report or args.input.with_suffix(".qa.json"), result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result["status"] == "pass" else 2
    except Exception as exc:
        print(f"fastgen: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
