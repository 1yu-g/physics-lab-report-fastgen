import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import fastgen
import workflow


class WorkflowIntegrationTest(unittest.TestCase):
    def test_prepare_build_review_gate_and_template_preservation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            template = root / "source.docx"
            guide = root / "guide.txt"
            figure = root / "diagram.png"
            work = root / "work"
            guide.write_text("一、实验目的\n记录真实数据。图1 原理示意图。", encoding="utf-8")
            Image.new("RGB", (300, 160), "white").save(figure)
            doc = Document()
            doc.styles["Normal"].font.name = "Arial"
            cover = doc.add_paragraph()
            cover.add_run("姓名：")
            cover.add_run("{{姓").bold = True
            cover.add_run("名}}").italic = True
            doc.add_paragraph("一、实验目的", style="Heading 1")
            doc.add_paragraph("二、实验原理", style="Heading 1")
            doc.save(template)
            original_hash = fastgen.sha256(template)

            result = workflow.prepare(
                work, [guide], template, [figure], [],
                ["一、实验目的", "二、实验原理"], "仅前两项，封面不动"
            )
            self.assertEqual(result["files"], 3)
            inventory = fastgen.load(work / "inventory.json")
            self.assertIn("{{姓名}}", inventory["files"][1]["placeholders"])
            self.assertEqual(inventory["requested_sections"], ["一、实验目的", "二、实验原理"])
            with self.assertRaises(ValueError):
                workflow.run(work, renderer="none")

            spec = fastgen.load(work / "report.json")
            spec["metadata"] = {"姓名": "示例"}
            spec["sections"][0]["blocks"] = [{"type": "paragraph", "text": "掌握测量原理。"}]
            spec["sections"][1]["blocks"] = [
                {"type": "formula", "text": "U_H = K_H I B"},
                {"type": "figure", "path": str(figure), "caption": "原理图"},
                {"type": "table", "caption": "记录表", "columns": ["序号", "读数"], "rows": [["1", ""]]},
            ]
            fastgen.dump(work / "report.json", spec)
            result = workflow.run(work, renderer="none")
            self.assertEqual(result["status"], "render-unavailable")
            built = Document(result["docx"])
            self.assertEqual(built.styles["Normal"].font.name, "Arial")
            self.assertIn("姓名：示例", built.paragraphs[0].text)
            self.assertEqual(fastgen.sha256(template), original_hash)
            self.assertEqual(fastgen.load(result["qa"])["status"], "pass")
            with self.assertRaises(ValueError):
                workflow.finalize(work)

            if shutil.which("pdftoppm"):
                pdf = root / "exported.pdf"
                writer = PdfWriter()
                writer.add_blank_page(width=595, height=842)
                with pdf.open("wb") as stream:
                    writer.write(stream)
                built_hash = fastgen.sha256(Path(result["docx"]))
                result = workflow.preview(work, pdf)
                self.assertEqual(fastgen.sha256(Path(result["docx"])), built_hash)
                self.assertEqual(result["status"], "needs-visual-review")
                self.assertEqual(result["pages"], 1)
                review = fastgen.load(work / "review.json")
                review["pages"][0]["status"] = "pass"
                fastgen.dump(work / "review.json", review)
                self.assertEqual(workflow.finalize(work)["pages_reviewed"], 1)


if __name__ == "__main__":
    unittest.main()
