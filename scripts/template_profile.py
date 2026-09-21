#!/usr/bin/env python3
"""Create and validate reusable structural profiles for fixed DOCX templates."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

import fastgen

PROFILE_VERSION = "1"


def inspect_template(template: Path):
    template = template.expanduser().resolve()
    if not template.is_file() or template.suffix.lower() != ".docx":
        raise ValueError(f"Template must be an existing DOCX: {template}")
    doc = Document(template)
    paragraphs = list(fastgen.paragraphs(doc))
    text = "\n".join(paragraph.text for paragraph in paragraphs)
    sections = []
    for section in doc.sections:
        sections.append({
            "width_emu": section.page_width,
            "height_emu": section.page_height,
            "top_margin_emu": section.top_margin,
            "bottom_margin_emu": section.bottom_margin,
            "left_margin_emu": section.left_margin,
            "right_margin_emu": section.right_margin,
        })
    return {
        "version": PROFILE_VERSION,
        "template": str(template),
        "template_sha256": fastgen.sha256(template),
        "paragraphs": len(paragraphs),
        "anchors": [
            {"index": index, "text": paragraph.text.strip(), "style": paragraph.style.name}
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip()
        ],
        "placeholders": sorted(set(fastgen.PLACEHOLDER.findall(text))),
        "tables": [
            {
                "index": index,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "preview": [cell.text.strip() for cell in table.rows[0].cells] if table.rows else [],
            }
            for index, table in enumerate(doc.tables)
        ],
        "inline_images": len(doc.inline_shapes),
        "native_math_objects": len(doc.element.xpath(".//m:oMath")),
        "sections": sections,
    }


def create(template: Path, output: Path):
    profile = inspect_template(template)
    output = output.expanduser().resolve()
    fastgen.dump(output, profile)
    return {"status": "pass", "profile": str(output), **profile}


def check(profile_path: Path, template: Path):
    profile_path = profile_path.expanduser().resolve()
    profile = fastgen.load(profile_path)
    if profile.get("version") != PROFILE_VERSION:
        raise ValueError("Template profile version is unsupported; create it again.")
    template = template.expanduser().resolve()
    if not template.is_file() or template.suffix.lower() != ".docx":
        raise ValueError(f"Template must be an existing DOCX: {template}")
    current_hash = fastgen.sha256(template)
    if profile.get("template_sha256") != current_hash:
        raise ValueError("Template changed after profiling; create a new profile.")
    return {
        "status": "pass", "profile": str(profile_path),
        "template": str(template), "template_sha256": current_hash,
        "anchors": len(profile.get("anchors", [])), "tables": len(profile.get("tables", [])),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    create_parser = sub.add_parser("create")
    create_parser.add_argument("--template", type=Path, required=True)
    create_parser.add_argument("--output", type=Path, required=True)
    check_parser = sub.add_parser("check")
    check_parser.add_argument("--profile", type=Path, required=True)
    check_parser.add_argument("--template", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = (create(args.template, args.output) if args.command == "create"
                  else check(args.profile, args.template))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f"template_profile: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
